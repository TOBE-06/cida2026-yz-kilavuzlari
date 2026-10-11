"""Eleme formu (Excel) üretimi.

Açılır listeler ve "Yönerge" sekmesi `config/eleme_olcutleri.yaml` dosyasından üretilir;
ölçüt metni bu modülde yer almaz. Araştırmacı sütunları betik tarafından doldurulmaz
(`ulke` ve `bolge` evrenden öneri olarak gelir, değiştirilebilir).
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BILGI_SUTUNLARI = [
    "aday_id", "on_puan", "kurum", "olasi_grup", "baslik", "url", "isaretler", "mevcut_belge", "bulunma_yolu",
    "dil", "yz_terim_sayisi", "normatif_isaret_sayisi", "metin_uzunlugu", "erisim_durumu", "eslesen_terimler",
]
ARASTIRMACI_SUTUNLARI = [
    "karar", "dislama_kodu", "gerekce", "grup", "aktor_turu", "ulke", "bolge", "belge_bicimi",
    "belge_statusu", "kodlama_kapsami", "birlestir",
]
GENISLIKLER = {"kurum": 26, "baslik": 44, "url": 50, "isaretler": 48, "gerekce": 36, "kodlama_kapsami": 30,
               "eslesen_terimler": 30, "aktor_turu": 26, "belge_statusu": 30, "belge_bicimi": 26, "bolge": 18}
BASLIK_RENGI, ARASTIRMACI_RENGI = "D9E1F2", "FFF2CC"


def kodlu_etiketler(sozluk: dict[Any, str]) -> list[str]:
    """`{1: "Ajans"}` → `["1 – Ajans"]` (açılır listede kod ve etiket birlikte görünür)."""
    return [f"{kod} – {etiket}" for kod, etiket in sozluk.items()]


def liste_secenekleri(olcut: dict[str, Any]) -> dict[str, list[str]]:
    """Araştırmacı sütunlarının açılır liste seçeneklerini ölçüt dosyasından üretir."""
    s = olcut["siniflama"]
    return {
        "karar": list(olcut["karar"]),
        "dislama_kodu": list(olcut["dislama"]),
        "grup": list(s["grup"]),
        "aktor_turu": kodlu_etiketler(s["aktor_turu"]),
        "bolge": kodlu_etiketler(s["bolge"]),
        "belge_bicimi": kodlu_etiketler(s["belge_bicimi"]),
        "belge_statusu": kodlu_etiketler(s["belge_statusu"]),
    }


def bolge_secenegi(etiket: str, olcut: dict[str, Any]) -> str:
    """Evrendeki bölge adını (`Avrupa`) formdaki seçeneğe (`2 – Avrupa`) çevirir; eşleşmezse boş."""
    for kod, ad in olcut["siniflama"]["bolge"].items():
        if ad == etiket:
            return f"{kod} – {ad}"
    return ""


def onceki_kararlar(form: Path) -> dict[str, dict[str, Any]]:
    """Var olan formdaki araştırmacı girişlerini `aday_id → sütunlar` olarak okur."""
    if not form.exists():
        return {}
    tablo = pd.read_excel(form, sheet_name="Eleme", dtype=str, keep_default_na=False)
    sutunlar = [s for s in ARASTIRMACI_SUTUNLARI if s in tablo.columns]
    return {s["aday_id"]: {k: s[k] for k in sutunlar} for s in tablo.to_dict("records")}


def _baslik_bicimle(sayfa, sutunlar: list[str]) -> None:
    for i, ad in enumerate(sutunlar, start=1):
        hucre = sayfa.cell(row=1, column=i, value=ad)
        hucre.font = Font(bold=True)
        hucre.alignment = Alignment(wrap_text=True, vertical="top")
        hucre.fill = PatternFill("solid", start_color=ARASTIRMACI_RENGI if ad in ARASTIRMACI_SUTUNLARI else BASLIK_RENGI)
        sayfa.column_dimensions[get_column_letter(i)].width = GENISLIKLER.get(ad, 14)


def _listeler_sayfasi(kitap: Workbook, secenekler: dict[str, list[str]]) -> dict[str, str]:
    """Seçenekleri gizli bir sayfaya yazar; her sütun için başvuru aralığını döndürür."""
    sayfa = kitap.create_sheet("Listeler")
    araliklar = {}
    for i, (ad, degerler) in enumerate(secenekler.items(), start=1):
        harf = get_column_letter(i)
        sayfa.cell(row=1, column=i, value=ad)
        for j, deger in enumerate(degerler, start=2):
            sayfa.cell(row=j, column=i, value=deger)
        araliklar[ad] = f"Listeler!${harf}$2:${harf}${len(degerler) + 1}"
    sayfa.sheet_state = "hidden"
    return araliklar


def _yonerge_sayfasi(kitap: Workbook, olcut: dict[str, Any], isaret_aciklamasi: list[str]) -> None:
    sayfa = kitap.create_sheet("Yönerge")
    sayfa.column_dimensions["A"].width, sayfa.column_dimensions["B"].width = 34, 120
    satirlar: list[tuple[str, str, bool]] = [("ELEME YÖNERGESİ", "Kaynak: config/eleme_olcutleri.yaml. Kararı araştırmacı verir.", True)]

    def bolum(ad: str, ogeler: list[tuple[str, str]]) -> None:
        satirlar.append(("", "", False))
        satirlar.append((ad, "", True))
        satirlar.extend((a, b, False) for a, b in ogeler)

    bolum("Dahil etme ölçütleri (hepsi sağlanmalı)", list(olcut["dahil_etme"].items()))
    bolum("Dışlama kodları (ilk eşleşen kod yazılır)", list(olcut["dislama"].items()))
    bolum("Eleme sırası", [
        (f"{i}. {adim['asama']}", ", ".join(adim.get("kodlar", [])) or "Doldurulacak: " + ", ".join(adim["sonra_doldurulacak"]))
        for i, adim in enumerate(olcut["eleme_sirasi"], start=1)
    ])
    for ad, degerler in olcut["siniflama"].items():
        bolum(f"Sınıflama: {ad}", [(str(k), str(v)) for k, v in degerler.items()])
    bolum("Sınır durumlar için karar kuralları", [(d["durum"], d["kural"]) for d in olcut["sinir_durumlar"]])
    bolum("Sütunlar hakkında", [(a, b) for a, _, b in (s.partition(": ") for s in isaret_aciklamasi)])
    for i, (a, b, kalin) in enumerate(satirlar, start=1):
        for j, deger in enumerate((a, b), start=1):
            hucre = sayfa.cell(row=i, column=j, value=deger)
            hucre.alignment = Alignment(wrap_text=True, vertical="top")
            hucre.font = Font(bold=kalin)


def form_yaz(tablo: pd.DataFrame, form: Path, olcut: dict[str, Any], isaret_aciklamasi: list[str]) -> None:
    """Eleme formunu yazar. `tablo` bilgi ve araştırmacı sütunlarını içermelidir."""
    sutunlar = BILGI_SUTUNLARI + ARASTIRMACI_SUTUNLARI
    kitap = Workbook()
    sayfa = kitap.active
    sayfa.title = "Eleme"
    _baslik_bicimle(sayfa, sutunlar)
    for r, satir in enumerate(tablo[sutunlar].to_dict("records"), start=2):
        for c, ad in enumerate(sutunlar, start=1):
            deger = satir[ad]
            hucre = sayfa.cell(row=r, column=c, value=None if deger == "" else deger)
            if ad == "url" and deger:
                hucre.hyperlink, hucre.font = deger, Font(color="0563C1", underline="single")
            if ad in ("baslik", "isaretler", "gerekce", "kodlama_kapsami"):
                hucre.alignment = Alignment(wrap_text=True, vertical="top")
    son_satir = max(len(tablo) + 1, 2)
    for ad, aralik in _listeler_sayfasi(kitap, liste_secenekleri(olcut)).items():
        harf = get_column_letter(sutunlar.index(ad) + 1)
        dogrulama = DataValidation(type="list", formula1=f"={aralik}", allow_blank=True)
        dogrulama.error, dogrulama.errorTitle = "Listeden bir değer seçin.", "Geçersiz değer"
        sayfa.add_data_validation(dogrulama)
        dogrulama.add(f"{harf}2:{harf}{son_satir}")
    sayfa.freeze_panes = "D2"
    sayfa.auto_filter.ref = f"A1:{get_column_letter(len(sutunlar))}{son_satir}"
    _yonerge_sayfasi(kitap, olcut, isaret_aciklamasi)
    form.parent.mkdir(parents=True, exist_ok=True)
    kitap.save(form)


def form_oku(form: Path) -> pd.DataFrame:
    """Doldurulmuş formun "Eleme" sayfasını metin olarak okur (Aşama 4 ve sonrası için)."""
    load_workbook(form, read_only=True).close()   # dosya bozuksa burada anlaşılır
    return pd.read_excel(form, sheet_name="Eleme", dtype=str, keep_default_na=False)
