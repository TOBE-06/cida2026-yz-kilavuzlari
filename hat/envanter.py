"""Envanter ve manifest üretimi (Aşama 5; Aşama 6 uzlaştırma sonucunu ekleyip yeniden üretir).

Envanter, araştırmacının `Dahil` dediği belgeleri içerir: belge başına bir satır. Çok parçalı
belgelerde adresler, tarihler ve özetler parça sırasıyla ` ; ` ile birleştirilir. Tüm değerler
eleme formundan ve arşivdeki üst veri (JSON) dosyalarından okunur; hiçbiri elle yazılmaz.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from hat import form as f
from hat.ayarlar import KOK, yol
from hat.kayit import sha256_dosya

ENVANTER_SUTUNLARI = [
    "belge_id", "kurum", "grup", "aktor_turu", "ulke", "bolge", "belge_adi", "url", "son_url", "erisim_tarihi_utc",
    "dil", "belge_bicimi", "belge_statusu", "kodlama_kapsami", "bulunma_yolu", "toplama_yontemi", "karar",
    "dislama_kodu", "gerekce", "sha256_pdf", "sha256_txt", "wayback_url", "mevcut_arsivle_eslesme",
    "surum_degisikligi", "not",
]
MANIFEST_SUTUNLARI = ["dosya", "belge_id", "parca", "uzanti", "boyut_bayt", "sha256", "ust_veriyle_uyumlu"]


def kod(deger: str) -> str:
    """Formdaki `2 – Avrupa` seçeneğinden sayısal kodu (`2`) alır; boşsa boş döner."""
    return deger.split(" – ")[0].strip() if deger else ""


def birlestir(degerler: list[Any], tekil: bool = False) -> str:
    """Boş olmayan değerleri ` ; ` ile birleştirir; `tekil` ise yinelenenleri bir kez yazar."""
    dolu = [str(d) for d in degerler if d not in ("", None)]
    return " ; ".join(dict.fromkeys(dolu) if tekil else dolu)


def belge_satiri(belge_id: str, parcalar: pd.DataFrame, form: pd.DataFrame, uzlastirma: dict[str, dict]) -> dict[str, str]:
    """Bir belgenin envanter satırını, parçalarının form satırlarından ve üst verilerinden üretir."""
    ust = [json.loads((KOK / y).read_text(encoding="utf-8")) for y in parcalar["ust_veri"] if y]
    satirlar = form.loc[list(parcalar["aday_id"])]
    ana = satirlar.iloc[0]
    notlar = [f"{len(parcalar)} parça: " + ", ".join(belge_id + p for p in parcalar["parca"])] if len(parcalar) > 1 else []
    kontrol = [u["belge_id"] + u["parca"] for u in ust if u.get("gorsel_kontrol_gerekli")]
    notlar += ["görsel kontrol gerekli: " + ", ".join(kontrol)] if kontrol else []
    u = uzlastirma.get(belge_id, {})
    return {
        "belge_id": belge_id, "kurum": ana["kurum"], "grup": ana["grup"], "aktor_turu": kod(ana["aktor_turu"]),
        "ulke": ana["ulke"], "bolge": kod(ana["bolge"]),
        "belge_adi": next((x.get("baslik") for x in ust if x.get("baslik")), ana["baslik"]),
        "url": birlestir([x["istenen_url"] for x in ust]), "son_url": birlestir([x.get("son_url") for x in ust]),
        "erisim_tarihi_utc": birlestir([x["erisim_zamani_utc"] for x in ust]),
        "dil": birlestir([x["dil"] for x in ust], tekil=True),
        "belge_bicimi": kod(ana["belge_bicimi"]), "belge_statusu": kod(ana["belge_statusu"]),
        "kodlama_kapsami": birlestir(list(satirlar["kodlama_kapsami"])),
        "bulunma_yolu": birlestir(list(satirlar["bulunma_yolu"]), tekil=True),
        "toplama_yontemi": birlestir([x["toplama_yontemi"] for x in ust], tekil=True),
        "karar": ana["karar"], "dislama_kodu": ana["dislama_kodu"], "gerekce": ana["gerekce"],
        "sha256_pdf": birlestir([x["dosyalar"].get("pdf", {}).get("sha256") for x in ust]),
        "sha256_txt": birlestir([x["dosyalar"].get("txt", {}).get("sha256") for x in ust]),
        "wayback_url": birlestir([x.get("wayback_url") for x in ust]),
        "mevcut_arsivle_eslesme": u.get("durum", ""), "surum_degisikligi": u.get("surum_degisikligi", ""),
        "not": "; ".join(notlar),
    }


def envanter_olustur(ayarlar: dict[str, Any]) -> pd.DataFrame:
    """Eleme formu, arşiv durumu ve (varsa) uzlaştırma tablosundan envanteri üretir."""
    form = f.form_oku(yol("eleme") / ayarlar["puanlama"]["form_dosyasi"]).set_index("aday_id", drop=False)
    durum = pd.read_csv(yol("ara") / ayarlar["arsiv"]["durum_dosyasi"], dtype=str, keep_default_na=False)
    durum = durum[(durum["durum"] == "arşivlendi") & durum["aday_id"].isin(form[form["karar"] == "Dahil"]["aday_id"])]
    uzl_dosyasi = yol("rapor") / ayarlar["uzlastirma"]["cikti_dosyasi"]
    uzlastirma = (pd.read_csv(uzl_dosyasi, dtype=str, keep_default_na=False).set_index("belge_id").to_dict("index")
                  if uzl_dosyasi.exists() else {})
    satirlar = [belge_satiri(b, p.sort_values("parca"), form, uzlastirma) for b, p in durum.groupby("belge_id")]
    return pd.DataFrame(satirlar, columns=ENVANTER_SUTUNLARI)


def envanter_yaz(envanter: pd.DataFrame, ayarlar: dict[str, Any], olcut: dict[str, Any]) -> tuple[Path, Path]:
    """Envanteri CSV ve Excel olarak yazar. Excel'de sayısal kodların anlamı "Kod defteri" sekmesindedir."""
    klasor = yol("envanter")
    klasor.mkdir(parents=True, exist_ok=True)
    csv, xlsx = klasor / "envanter.csv", klasor / "envanter.xlsx"
    envanter.to_csv(csv, index=False, encoding="utf-8")
    defter = pd.DataFrame([{"degisken": ad, "kod": k, "etiket": v} for ad, d in olcut["siniflama"].items() for k, v in d.items()])
    with pd.ExcelWriter(xlsx, engine="openpyxl") as yazici:
        envanter.to_excel(yazici, sheet_name="Envanter", index=False)
        defter.to_excel(yazici, sheet_name="Kod defteri", index=False)
        sayfa = yazici.sheets["Envanter"]
        sayfa.freeze_panes = "C2"
        for i, ad in enumerate(envanter.columns, start=1):
            sayfa.cell(row=1, column=i).font = Font(bold=True)
            sayfa.column_dimensions[get_column_letter(i)].width = 14 if not ad.startswith(("kurum", "belge_adi", "url")) else 36
    return csv, xlsx


def manifest_olustur(ayarlar: dict[str, Any]) -> pd.DataFrame:
    """`archive/K` ve `archive/M` altındaki her dosyanın boyutunu ve SHA-256 özetini diskten hesaplar.

    Özet, üst verideki kayıtla karşılaştırılır; fark varsa `ust_veriyle_uyumlu = hayır` yazılır.
    """
    kayitli: dict[str, str] = {}
    dosyalar = sorted(d for g in ("K", "M") for d in (yol("arsiv") / g).glob("*") if d.is_file())
    for d in dosyalar:
        if d.suffix == ".json":
            kayitli |= {v["ad"]: v["sha256"] for v in json.loads(d.read_text(encoding="utf-8"))["dosyalar"].values()}
    satirlar = []
    for d in dosyalar:
        m = re.match(r"([KM]\d+)([a-z]?)_", d.name)
        ozet = sha256_dosya(d)
        uyum = "üst verinin kendisi" if d.suffix == ".json" else ("evet" if kayitli.get(d.name) == ozet else "hayır")
        satirlar.append({"dosya": str(d.relative_to(KOK)), "belge_id": m.group(1) if m else "", "parca": m.group(2) if m else "",
                         "uzanti": d.suffix.lstrip("."), "boyut_bayt": d.stat().st_size, "sha256": ozet, "ust_veriyle_uyumlu": uyum})
    return pd.DataFrame(satirlar, columns=MANIFEST_SUTUNLARI)
