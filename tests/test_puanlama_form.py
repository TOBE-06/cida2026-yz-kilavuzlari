"""Ön puanlama, işaretler ve eleme formu testleri."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from hat import form as f
from hat import puanlama as p
from hat.ayarlar import ayarlari_yukle, eleme_olcutleri
from hat.metin import html_metin

AYAR = ayarlari_yukle()["puanlama"]
OLCUT = eleme_olcutleri()
ORNEK = (Path(__file__).parent / "ornekler" / "kilavuz_en.html").read_text(encoding="utf-8")


def test_on_puan_tavanlari_ve_baslik_puani():
    assert p.on_puan(0, 0, False, AYAR) == 0
    assert p.on_puan(3, 5, True, AYAR) == 3 * 2 + 5 + 20
    assert p.on_puan(999, 999, True, AYAR) == 100


def test_metin_olcumleri_ornek_kilavuz():
    o = p.metin_olcumleri("https://ornek.com/about/ai-policy", "Our approach to generative AI", html_metin(ORNEK), ORNEK, AYAR)
    assert o["dil"] == "en" and o["yz_terim_sayisi"] >= 4 and o["normatif_isaret_sayisi"] >= 4
    assert o["on_puan"] > 20 and len(o["metin_ozeti"]) == 64


def test_benzer_adaylar():
    metinler = {"AD1": "yapay zekâ ilkeleri " * 50, "AD2": "yapay zekâ ilkeleri " * 50 + "ek", "AD3": "spor haberi " * 80, "AD4": ""}
    assert p.benzer_adaylar(metinler, 95, 20000) == {"AD1": ["AD2"], "AD2": ["AD1"]}


def _aday(**ek):
    return {"erisim_durumu": "tamam", "yz_terim_sayisi": "5", "normatif_isaret_sayisi": "3", "metin_uzunlugu": "2000"} | ek


def test_isaretler_gozlem_yazar_kod_onermez():
    metin = p.isaretler(_aday(yz_terim_sayisi="0", metin_uzunlugu="120", mevcut_belge="K07"), 5, ["AD0009"], True, 400)
    for beklenen in ("mevcut 47 belgeden biriyle eşleşiyor: K07", "dışlanan aktörler listesinde", "metinde YZ terimi yok",
                     "metin çok kısa", "aynı kurumdan 5 aday", "metni çok benzer aday: AD0009"):
        assert beklenen in metin
    assert not any(f"E{i}" in metin for i in range(1, 9))


def test_isaretler_erisilemeyen_sayfa():
    metin = p.isaretler(_aday(erisim_durumu="engellendi", erisim_aciklamasi="HTTP 403", wayback_url="https://web.archive.org/x"), 1, [], False, 400)
    assert "sayfaya erişilemedi (engellendi: HTTP 403)" in metin and "Wayback kopyası var" in metin
    assert "YZ terimi yok" not in metin
    assert p.isaretler(_aday(), 1, [], False, 400) == ""


def _tablo() -> pd.DataFrame:
    satir = {s: "" for s in f.BILGI_SUTUNLARI + f.ARASTIRMACI_SUTUNLARI} | {
        "aday_id": "AD0001", "on_puan": 55, "kurum": "Örnek", "url": "https://ornek.com/ai-policy", "ulke": "İngiltere",
        "bolge": f.bolge_secenegi("Avrupa", OLCUT),
    }
    return pd.DataFrame([satir, satir | {"aday_id": "AD0002", "on_puan": 10}])


def test_form_olcut_dosyasindan_uretilir(tmp_path):
    dosya = tmp_path / "form.xlsx"
    f.form_yaz(_tablo(), dosya, OLCUT, ["on_puan: karar değildir."])
    kitap = load_workbook(dosya)
    assert kitap.sheetnames == ["Eleme", "Listeler", "Yönerge"]
    eleme = kitap["Eleme"]
    basliklar = [h.value for h in eleme[1]]
    assert basliklar == f.BILGI_SUTUNLARI + f.ARASTIRMACI_SUTUNLARI and eleme.freeze_panes == "D2"
    assert eleme.cell(row=2, column=basliklar.index("url") + 1).hyperlink.target == "https://ornek.com/ai-policy"
    assert eleme.cell(row=2, column=basliklar.index("karar") + 1).value is None          # karar boş
    assert eleme.cell(row=2, column=basliklar.index("bolge") + 1).value == "2 – Avrupa"   # evrenden öneri
    assert len(eleme.data_validations.dataValidation) == 7
    listeler = {s[0].value: [h.value for h in s[1:] if h.value] for s in kitap["Listeler"].iter_cols()}
    assert listeler["karar"] == OLCUT["karar"] and listeler["dislama_kodu"] == list(OLCUT["dislama"])
    assert listeler["belge_statusu"][0] == "1 – " + OLCUT["siniflama"]["belge_statusu"][1]
    yonerge = " ".join(str(h.value) for s in kitap["Yönerge"].iter_rows() for h in s if h.value)
    assert OLCUT["dislama"]["E6"] in yonerge and OLCUT["sinir_durumlar"][0]["kural"] in yonerge


def test_onceki_kararlar_yeni_forma_tasinir(tmp_path):
    dosya = tmp_path / "form.xlsx"
    f.form_yaz(_tablo(), dosya, OLCUT, [])
    kitap = load_workbook(dosya)
    basliklar = [h.value for h in kitap["Eleme"][1]]
    kitap["Eleme"].cell(row=3, column=basliklar.index("karar") + 1, value="Hariç")
    kitap["Eleme"].cell(row=3, column=basliklar.index("dislama_kodu") + 1, value="E3")
    kitap.save(dosya)
    onceki = f.onceki_kararlar(dosya)
    assert onceki["AD0002"]["karar"] == "Hariç" and onceki["AD0002"]["dislama_kodu"] == "E3"
    assert onceki["AD0001"]["karar"] == ""
